import csv
import json
import math
import os
import shutil
import sys
import threading
import traceback
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
import tkinter as tk

import torch
import torch.nn as nn
from PIL import Image, ImageTk
from torchvision import transforms
from torchvision.models import efficientnet_b3, mobilenet_v2, resnet18


APP_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
RUNTIME_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
MODELS_DIR = RUNTIME_DIR / "models"
FEATURES_DIR = RUNTIME_DIR / "features"
TRAIN_STATS_PATH = FEATURES_DIR / "train_stats.json"
if not getattr(sys, "frozen", False):
    from workspace_paths import default_path
    MODELS_DIR = default_path('models', MODELS_DIR)
    FEATURES_DIR = default_path('reference_features', FEATURES_DIR)
# train_stats remains the original packaged legacy resource, not another model's stats.
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".jfif"}
PROCESS_ORDER = ("convert", "crop", "rb", "assort")
PROCESS_STEPS = {
    "convert": {
        "label": "Format convert (JFIF -> JPG)",
        "warning": "This step deletes .jfif files only inside the IMG_result working copy.",
    },
    "crop": {"label": "Batch crop"},
    "rb": {"label": "Blue-line removal / RB"},
    "assort": {"label": "Image classification split"},
}


def is_path_sequence(path):
    return isinstance(path, (list, tuple, set))

CLASS_NAMES = {
    0: "eGFR -15",
    1: "eGFR 15-29",
    2: "eGFR 30-59",
    3: "eGFR 60-89",
    4: "eGFR 90-",
}
CLASS_COUNT = 5
META_DIM = 4
META_COLS = ["age", "Length_px", "Contrast_Ratio", "has_diabetes"]
BINARY_META_COLS = {"gender", "has_diabetes", "diabetes"}
FEATURE_OUTPUT_FIELDS = ["image_name", "age", "gender", "Length_px", "Entropy", "Contrast_Ratio", "has_diabetes"]


class ResNetBackbone(nn.Module):
    def __init__(self):
        super().__init__()
        backbone = resnet18(weights=None)
        self.features = nn.Sequential(*list(backbone.children())[:-1])
        self.feature_dim = 512

    def forward(self, x):
        feat = self.features(x)
        return feat.view(feat.size(0), -1)


class MobileNetV2Backbone(nn.Module):
    def __init__(self):
        super().__init__()
        backbone = mobilenet_v2(weights=None)
        self.features = backbone.features
        self.avgpool = nn.AdaptiveAvgPool2d((1, 1))
        self.feature_dim = 1280

    def forward(self, x):
        x = self.features(x)
        x = self.avgpool(x)
        return x.view(x.size(0), -1)


class EfficientNetB3Backbone(nn.Module):
    def __init__(self):
        super().__init__()
        backbone = efficientnet_b3(weights=None)
        self.features = backbone.features
        self.pool = nn.AdaptiveAvgPool2d(1)
        self.feature_dim = backbone.classifier[1].in_features

    def forward(self, x):
        x = self.features(x)
        x = self.pool(x)
        return torch.flatten(x, 1)


def make_backbone(backbone_type):
    if backbone_type == "resnet18":
        return ResNetBackbone()
    if backbone_type == "mobilenetv2":
        return MobileNetV2Backbone()
    if backbone_type == "efficientnet_b3":
        return EfficientNetB3Backbone()
    raise ValueError(f"Unsupported backbone: {backbone_type}")


class BasicResNetClassifier(nn.Module):
    def __init__(self, num_classes=CLASS_COUNT, backbone_type="resnet18"):
        super().__init__()
        self.num_classes = num_classes
        self.backbone = make_backbone(backbone_type)
        in_features = self.backbone.feature_dim
        out_features = 1 if num_classes == 2 else num_classes
        self.classifier = nn.Sequential(
            nn.Linear(in_features, 256),
            nn.ReLU(),
            nn.BatchNorm1d(256),
            nn.Dropout(0.5),
            nn.Linear(256, out_features),
        )

    def forward(self, x):
        logits = self.classifier(self.backbone(x))
        return logits.squeeze(-1) if self.num_classes == 2 else logits


class MLPClassifier(nn.Module):
    def __init__(self, num_classes=CLASS_COUNT, meta_dim=META_DIM, backbone_type="resnet18"):
        super().__init__()
        self.num_classes = num_classes
        self.meta_dim = meta_dim
        self.backbone = make_backbone(backbone_type)
        in_features = self.backbone.feature_dim
        self.feature_mlp = nn.Sequential(
            nn.Linear(in_features, 256),
            nn.ReLU(),
            nn.BatchNorm1d(256),
            nn.Dropout(0.5),
        )
        self.meta_mlp = nn.Sequential(
            nn.Linear(meta_dim, 32),
            nn.ReLU(),
            nn.BatchNorm1d(32),
            nn.Dropout(0.3),
        )
        out_features = 1 if num_classes == 2 else num_classes
        self.classifier = nn.Linear(256 + 32, out_features)

    def forward(self, x, meta):
        feat = self.feature_mlp(self.backbone(x))
        meta_feat = self.meta_mlp(meta)
        logits = self.classifier(torch.cat([feat, meta_feat], dim=1))
        return logits.squeeze(-1) if self.num_classes == 2 else logits


def load_model_weights(model, path, device):
    try:
        state = torch.load(path, map_location=device, weights_only=True)
    except TypeError:
        state = torch.load(path, map_location=device)
    model.load_state_dict(state)
    return model


def list_images(path):
    if is_path_sequence(path):
        images = []
        for item in path:
            images.extend(list_images(item))
        return sorted(dict.fromkeys(images))
    path = Path(path)
    if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES:
        return [path]
    if path.is_dir():
        return sorted(
            p for p in path.rglob("*")
            if p.is_file()
            and p.suffix.lower() in IMAGE_SUFFIXES
            and "IMG_result" not in p.parts
        )
    return []


def image_root_for(path):
    if is_path_sequence(path):
        images = list_images(path)
        return images[0].parent if images else RUNTIME_DIR
    path = Path(path)
    return path.parent if path.is_file() else path


def result_parent_for(path):
    if is_path_sequence(path):
        images = list_images(path)
        if not images:
            return RUNTIME_DIR
        parents = [str(p.parent) for p in images]
        return Path(os.path.commonpath(parents))
    path = Path(path)
    return path.parent if path.is_file() else path.parent


def make_result_run(input_path):
    result_root = result_parent_for(input_path) / "IMG_result"
    result_root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("run_%Y%m%d_%H%M%S")
    run_dir = result_root / stamp
    suffix = 1
    while run_dir.exists():
        suffix += 1
        run_dir = result_root / f"{stamp}_{suffix}"
    run_dir.mkdir(parents=True)
    return run_dir


def copy_images_to_workspace(input_path, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    images = list_images(input_path)
    if not images:
        raise ValueError(f"No images found: {input_path}")
    if is_path_sequence(input_path):
        base = Path(os.path.commonpath([str(p.parent) for p in images]))
    else:
        input_path = Path(input_path)
        base = input_path.parent if input_path.is_file() else input_path
    copied = []
    for image_path in images:
        relative = image_path.name if not is_path_sequence(input_path) and Path(input_path).is_file() else image_path.relative_to(base)
        target = output_dir / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(image_path, target)
        copied.append(target)
    return output_dir, copied


def ensure_stdio_for_windowed_app():
    if sys.stdout is None:
        sys.stdout = open(os.devnull, "w", encoding="utf-8")
    if sys.stderr is None:
        sys.stderr = open(os.devnull, "w", encoding="utf-8")


def safe_float(value, default=0.0):
    if value is None or value == "":
        return default
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default
    return default if math.isnan(number) else number


def read_csv_rows(path):
    path = Path(path)
    if not path.exists():
        return []
    for encoding in ("utf-8-sig", "utf-8", "cp950"):
        try:
            with open(path, newline="", encoding=encoding) as f:
                return list(csv.DictReader(f))
        except UnicodeDecodeError:
            continue
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        return list(csv.DictReader(f))


def load_metadata_by_name(path):
    rows = read_csv_rows(path)
    metadata = {}
    for row in rows:
        name = row.get("image_name") or row.get("file") or row.get("filename")
        if name:
            metadata[Path(name).name] = row
    return metadata


def load_train_stats():
    if TRAIN_STATS_PATH.exists():
        with open(TRAIN_STATS_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {
        "age": {"mean": 0.0, "std": 1.0},
        "Length_px": {"mean": 0.0, "std": 1.0},
        "Contrast_Ratio": {"mean": 0.0, "std": 1.0},
    }


def normalize_meta_row(row, stats):
    values = []
    defaults = {
        "age": stats.get("age", {}).get("mean", 0.0),
        "Length_px": stats.get("Length_px", {}).get("mean", 0.0),
        "Contrast_Ratio": stats.get("Contrast_Ratio", {}).get("mean", 1.0),
        "has_diabetes": 0.0,
    }
    for col in META_COLS:
        raw_value = safe_float(row.get(col), defaults.get(col, 0.0))
        if col in BINARY_META_COLS:
            values.append(float(raw_value))
        else:
            stat = stats.get(col, {"mean": 0.0, "std": 1.0})
            values.append((raw_value - stat.get("mean", 0.0)) / (stat.get("std", 1.0) + 1e-10))
    return values


def build_metadata_tensors(csv_path, stats):
    metadata = {}
    for name, row in load_metadata_by_name(csv_path).items():
        metadata[name] = normalize_meta_row(row, stats)
    return metadata


def index_image_files(folder):
    files = {}
    for path in list_images(folder):
        files[path.stem] = path
        if path.stem.endswith("_mask"):
            files[path.stem[:-5]] = path
    return files


def infer_mask_dir(image_root):
    root = image_root_for(image_root)
    candidates = [
        root.parent / "IMG_Mask",
        root / "IMG_Mask",
        root.parent.parent / "IMG_Mask",
    ]
    for candidate in candidates:
        if candidate.exists() and list_images(candidate):
            return candidate
    return candidates[0]


def load_excel_diabetes_map(excel_path):
    excel_path = Path(excel_path)
    if not excel_path.exists():
        return []
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    try:
        egfr_index = headers.index("matched_eGFR")
    except ValueError:
        egfr_index = 0
    diabetes_index = None
    for index, header in enumerate(headers):
        lowered = header.lower()
        if "diabetes" in lowered or "dm" == lowered or "糖尿" in header:
            diabetes_index = index
            break
    if diabetes_index is None and len(headers) > 1:
        diabetes_index = len(headers) - 1
    if diabetes_index is None:
        return []

    mapping = []
    for row in rows[1:]:
        if len(row) <= max(egfr_index, diabetes_index or 0):
            continue
        egfr = safe_float(row[egfr_index], None)
        if egfr is None:
            continue
        mapping.append((egfr, safe_float(row[diabetes_index], 0.0)))
    return mapping


def parse_egfr_from_name(name):
    try:
        return float(Path(name).stem.split("_")[0])
    except (ValueError, IndexError):
        return None


def match_diabetes_from_egfr(image_name, diabetes_map, tolerance=0.2):
    egfr = parse_egfr_from_name(image_name)
    if egfr is None or not diabetes_map:
        return 0
    best = min(diabetes_map, key=lambda item: abs(item[0] - egfr))
    return int(best[1]) if abs(best[0] - egfr) <= tolerance else 0


def calculate_image_features(image_path, mask_path):
    try:
        import cv2
        import numpy as np
        from scipy.ndimage import find_objects
        from skimage import measure
        from skimage.feature import graycomatrix
    except ImportError as exc:
        raise RuntimeError(
            "Feature extraction needs cv2, numpy, scipy, and scikit-image. "
            "Install the same packages used by features/addmore.py."
        ) from exc

    config = {
        "glcm_levels": 32,
        "glcm_distances": [1, 2],
        "glcm_angles": [0, np.pi / 4, np.pi / 2, 3 * np.pi / 4],
        "contrast_kernel_min": 20,
        "contrast_kernel_max": 50,
        "contrast_kernel_scale": 0.35,
        "min_roi_pixels": 100,
        "denoise_kernel": 3,
        "epsilon": 1e-10,
    }

    img = cv2.imread(str(image_path))
    mask = cv2.imread(str(mask_path), cv2.IMREAD_GRAYSCALE)
    if img is None or mask is None:
        raise ValueError(f"Cannot read image or mask: {image_path.name}")
    if mask.shape != img.shape[:2]:
        mask = cv2.resize(mask, (img.shape[1], img.shape[0]), interpolation=cv2.INTER_NEAREST)
    _, mask_binary = cv2.threshold(mask, 127, 1, cv2.THRESH_BINARY)

    if np.sum(mask_binary) < config["min_roi_pixels"]:
        return {"Length_px": 0.0, "Entropy": 0.0, "Contrast_Ratio": 0.0}

    labeled_mask = measure.label(mask_binary)
    regions = measure.regionprops(labeled_mask)
    region = max(regions, key=lambda r: r.area) if regions else None
    length = float(region.major_axis_length) if region else 0.0

    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    if config["denoise_kernel"] > 0:
        gray = cv2.medianBlur(gray, config["denoise_kernel"])
    roi_pixels = gray[mask_binary == 1]
    mean_intensity = float(np.mean(roi_pixels)) if roi_pixels.size else 0.0

    entropy_mean = 0.0
    if find_objects(labeled_mask) and region:
        bbox = region.bbox
        cropped_gray = gray[bbox[0]:bbox[2], bbox[1]:bbox[3]]
        cropped_mask = mask_binary[bbox[0]:bbox[2], bbox[1]:bbox[3]]
        gray_masked = cropped_gray.copy()
        gray_masked[cropped_mask == 0] = 0
        max_val = np.max(gray_masked)
        if max_val > 0:
            levels = config["glcm_levels"]
            gray_scaled = (gray_masked / max_val * (levels - 1)).astype(np.uint8)
            glcm = graycomatrix(
                gray_scaled,
                distances=config["glcm_distances"],
                angles=config["glcm_angles"],
                levels=levels,
                symmetric=True,
                normed=True,
            )
            glcm_sum = glcm.sum(axis=(0, 1), keepdims=True)
            probs = glcm / (glcm_sum + config["epsilon"])
            entropy_mean = float(np.mean(-np.sum(probs * np.log2(probs + config["epsilon"]), axis=(0, 1))))

    roi_area = np.sum(mask_binary)
    kernel_size = int(np.sqrt(roi_area) * config["contrast_kernel_scale"])
    kernel_size = int(np.clip(kernel_size, config["contrast_kernel_min"], config["contrast_kernel_max"]))
    kernel_size = kernel_size if kernel_size % 2 == 1 else kernel_size + 1
    kernel = np.ones((kernel_size, kernel_size), np.uint8)
    dilated_mask = cv2.dilate(mask_binary.astype(np.uint8), kernel, iterations=1)
    ring_mask = np.clip(dilated_mask - mask_binary, 0, 1)
    background_pixels = gray[ring_mask == 1]
    if background_pixels.size < 10:
        contrast_ratio = 1.0
    else:
        contrast_ratio = float(mean_intensity / (np.mean(background_pixels) + config["epsilon"]))

    return {"Length_px": length, "Entropy": entropy_mean, "Contrast_Ratio": contrast_ratio}


def generate_feature_csv(images, output_csv, mask_dir=None, base_metadata_csv=None, diabetes_excel=None, progress=None):
    images = sorted(Path(p) for p in images)
    if not images:
        raise ValueError("No images found for feature extraction.")
    mask_dir = Path(mask_dir) if mask_dir else infer_mask_dir(image_root_for(images[0]))
    mask_index = index_image_files(mask_dir)
    base_metadata = load_metadata_by_name(base_metadata_csv) if base_metadata_csv else {}
    diabetes_map = load_excel_diabetes_map(diabetes_excel or (FEATURES_DIR / "add_HbA1C.xlsx"))
    rows = []
    missing_masks = 0

    for index, image_path in enumerate(images, 1):
        if progress:
            progress(index - 1, len(images), f"Calculating features {index}/{len(images)}")
        mask_path = mask_index.get(image_path.stem)
        base_row = base_metadata.get(image_path.name, {})
        row = {field: "" for field in FEATURE_OUTPUT_FIELDS}
        row.update({k: v for k, v in base_row.items() if k in row})
        row["image_name"] = image_path.name
        row["age"] = base_row.get("age", "")
        row["gender"] = base_row.get("gender", "")
        row["has_diabetes"] = base_row.get("has_diabetes", "")
        if row["has_diabetes"] == "":
            row["has_diabetes"] = match_diabetes_from_egfr(image_path.name, diabetes_map)

        if mask_path:
            row.update(calculate_image_features(image_path, mask_path))
        else:
            missing_masks += 1
            row.update({"Length_px": "", "Entropy": "", "Contrast_Ratio": ""})
        rows.append(row)

    output_csv = Path(output_csv)
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with open(output_csv, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=FEATURE_OUTPUT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    if progress:
        progress(len(images), len(images), f"Feature CSV exported: {output_csv}")
    return output_csv, len(rows), missing_masks


class InferenceEngine:
    def __init__(self, model_path, architecture, device=None, metadata=None, stats=None):
        self.model_path = Path(model_path)
        self.architecture = architecture
        self.device = device or torch.device("cpu")
        self.metadata = metadata or {}
        self.stats = stats or load_train_stats()
        self.transform = transforms.Compose(
            [
                transforms.Resize((224, 224)),
                transforms.ToTensor(),
                transforms.Normalize(mean=[0.485, 0.456, 0.406], std=[0.229, 0.224, 0.225]),
            ]
        )
        self.model = self._load_model().to(self.device)
        self.model.eval()

    def _load_model(self):
        arch = self.architecture.lower()
        if arch == "resnet18":
            model = BasicResNetClassifier(num_classes=CLASS_COUNT, backbone_type="resnet18")
        elif arch == "mobilenetv2":
            model = BasicResNetClassifier(num_classes=CLASS_COUNT, backbone_type="mobilenetv2")
        elif arch == "efficientnetb3":
            model = BasicResNetClassifier(num_classes=CLASS_COUNT, backbone_type="efficientnet_b3")
        elif arch == "mlp-resnet18":
            model = MLPClassifier(num_classes=CLASS_COUNT, meta_dim=META_DIM, backbone_type="resnet18")
        else:
            raise ValueError(f"Unsupported architecture: {self.architecture}")
        return load_model_weights(model, str(self.model_path), self.device)

    def _meta_tensor_for(self, image_path):
        values = self.metadata.get(Path(image_path).name)
        if values is None:
            values = normalize_meta_row({}, self.stats)
        return torch.tensor(values, dtype=torch.float32, device=self.device).unsqueeze(0)

    def predict(self, image_path):
        image = Image.open(image_path).convert("RGB")
        tensor = self.transform(image).unsqueeze(0).to(self.device)
        with torch.no_grad():
            if self.architecture.lower() == "mlp-resnet18":
                meta = self._meta_tensor_for(image_path)
                output = self.model(tensor, meta)
            else:
                output = self.model(tensor)

            if CLASS_COUNT == 2:
                p1 = torch.sigmoid(output).reshape(1)
                probs = torch.stack([1 - p1, p1], dim=1)[0]
            else:
                probs = torch.softmax(output, dim=1)[0]

        pred_index = int(torch.argmax(probs).item())
        confidence = float(probs[pred_index].item())
        return pred_index, CLASS_NAMES.get(pred_index, str(pred_index)), confidence, [float(p) for p in probs.cpu()]


class DeployApp:
    def __init__(self, root):
        self.root = root
        self.root.title("CNN Image Processing and Inference")
        self.root.geometry("1080x720")
        self.root.minsize(960, 620)
        self.selected_input = None
        self.selected_mask_input = None
        self.last_result_run = None
        self.results = []
        self.preview_image = None
        self.processing_busy = False
        self.mask_busy = False
        self.inference_busy = False
        self.model_var = tk.StringVar()
        self.unet_model_var = tk.StringVar()
        self.arch_var = tk.StringVar(value="ResNet18")
        self.status_var = tk.StringVar(value="Ready")
        self.progress_var = tk.DoubleVar(value=0)
        self.mask_output_var = tk.StringVar(value="")
        self.feature_csv_var = tk.StringVar(value="")
        self.mask_folder_var = tk.StringVar(value="")
        self.process_mask_folder_var = tk.StringVar(value="")
        self.process_step_vars = {key: tk.BooleanVar(value=False) for key in PROCESS_ORDER}
        self.process_input_vars = {key: tk.StringVar(value="") for key in PROCESS_ORDER}
        self.custom_process_order = list(PROCESS_ORDER)
        self.process_order_list = None
        self.process_mask_widgets = []
        self.process_step_widgets = {}
        self.process_override_widgets = {}
        self.process_move_buttons = []
        self.run_process_button = None
        self.infer_control_widgets = []
        self.infer_feature_widgets = []
        self.feature_csv_widgets = []
        self.run_feature_csv_button = None
        self.feature_summary = None
        self.run_infer_button = None
        self.export_button = None
        self.mask_model_widgets = []
        self.unet_model_combo = None
        self.run_mask_button = None
        self.nav_buttons = {}
        self.current_page = None

        self._build_layout()
        self.refresh_models()
        self._refresh_all_workflow_controls()

    def _build_layout(self):
        self.root.columnconfigure(1, weight=1)
        self.root.rowconfigure(0, weight=1)

        nav = ttk.Frame(self.root, padding=16)
        nav.grid(row=0, column=0, sticky="ns")
        ttk.Label(nav, text="Main Menu", font=("Segoe UI", 18, "bold")).pack(anchor="w", pady=(0, 18))
        nav_items = (
            ("process", "Image Processing"),
            ("mask", "Generate Mask"),
            ("feature_csv", "Feature CSV"),
            ("infer", "Inference & Report"),
            ("settings", "Settings"),
        )
        for page_name, label in nav_items:
            button = tk.Button(
                nav,
                text=label,
                command=lambda name=page_name: self.show_page(name),
                anchor="w",
                padx=12,
                pady=7,
                relief=tk.FLAT,
                borderwidth=0,
                highlightthickness=0,
            )
            button.pack(fill="x", pady=5)
            self.nav_buttons[page_name] = button

        content = ttk.Frame(self.root, padding=18)
        content.grid(row=0, column=1, sticky="nsew")
        content.columnconfigure(0, weight=1)
        content.rowconfigure(0, weight=1)

        self.pages = {
            "process": self._create_process_page(content),
            "mask": self._create_mask_page(content),
            "feature_csv": self._create_feature_csv_page(content),
            "infer": self._create_infer_page(content),
            "settings": self._create_settings_page(content),
        }
        self.show_page("process")

        bottom = ttk.Frame(self.root, padding=(18, 0, 18, 14))
        bottom.grid(row=1, column=0, columnspan=2, sticky="ew")
        bottom.columnconfigure(0, weight=1)
        ttk.Label(bottom, textvariable=self.status_var).grid(row=0, column=0, sticky="w")
        ttk.Progressbar(bottom, variable=self.progress_var, maximum=100).grid(row=0, column=1, sticky="ew", padx=(12, 0))

    def _create_process_page(self, parent):
        page = ttk.Frame(parent)
        page.columnconfigure(0, weight=1)
        ttk.Label(page, text="Image Processing", font=("Segoe UI", 20, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(
            page,
            #text="Order: 1 Select image/folder -> 2 Choose standard/custom and steps -> 3 Sort or set inputs -> 4 Run.",
        ).grid(row=1, column=0, sticky="w", pady=(4, 16))

        input_frame = ttk.LabelFrame(page, text="Step 1 - Input", padding=10)
        input_frame.grid(row=2, column=0, sticky="ew", pady=(0, 10))
        input_frame.columnconfigure(1, weight=1)
        input_frame.columnconfigure(3, weight=1)

        image_actions = ttk.Frame(input_frame)
        image_actions.grid(row=0, column=0, sticky="w", padx=(0, 8))
        ttk.Button(image_actions, text="Select Image", command=self.select_image).pack(side="left", padx=(0, 8))
        ttk.Button(image_actions, text="Select Folder", command=self.select_folder).pack(side="left")
        ttk.Label(input_frame, text="Images").grid(row=0, column=1, sticky="w", padx=(4, 6))
        self.process_summary = ttk.Label(input_frame, text="No input selected.")
        self.process_summary.grid(row=0, column=2, columnspan=3, sticky="w", padx=(0, 0))

        ttk.Label(input_frame, text="Mask folder").grid(row=1, column=0, sticky="w", pady=(8, 0))
        mask_entry = ttk.Entry(input_frame, textvariable=self.process_mask_folder_var)
        mask_entry.grid(row=1, column=1, columnspan=2, sticky="ew", pady=(8, 0), padx=(4, 6))
        mask_browse = ttk.Button(input_frame, text="Browse", command=self._browse_process_mask_folder)
        mask_browse.grid(row=1, column=3, sticky="ew", pady=(8, 0), padx=(0, 6))
        mask_clear = ttk.Button(input_frame, text="Clear", command=self._clear_process_mask_folder)
        mask_clear.grid(row=1, column=4, sticky="ew", pady=(8, 0))
        self.process_mask_widgets.extend([mask_entry, mask_browse, mask_clear])

        body = ttk.Frame(page)
        body.grid(row=3, column=0, sticky="nsew")
        body.columnconfigure(0, weight=1)
        body.columnconfigure(1, weight=1)
        page.rowconfigure(3, weight=1)

        steps_frame = ttk.LabelFrame(body, text="Step 2 - Custom processing steps", padding=10)
        steps_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        for key in PROCESS_ORDER:
            checkbox = ttk.Checkbutton(
                steps_frame,
                text=PROCESS_STEPS[key]["label"],
                variable=self.process_step_vars[key],
                command=self._refresh_process_order_ui,
            )
            checkbox.pack(anchor="w", pady=3)
            self.process_step_widgets[key] = checkbox

        order_frame = ttk.LabelFrame(body, text="Step 2 - Order and optional input folders", padding=10)
        order_frame.grid(row=0, column=1, sticky="nsew")
        order_frame.columnconfigure(0, weight=1)
        self.process_order_list = tk.Listbox(order_frame, height=7, exportselection=False)
        self.process_order_list.grid(row=0, column=0, rowspan=3, sticky="nsew", padx=(0, 8))
        move_up = ttk.Button(order_frame, text="Move up", command=lambda: self._move_process_step(-1))
        move_up.grid(row=0, column=1, sticky="ew", pady=(0, 6))
        move_down = ttk.Button(order_frame, text="Move down", command=lambda: self._move_process_step(1))
        move_down.grid(row=1, column=1, sticky="ew")
        self.process_move_buttons.extend([move_up, move_down])

        input_frame = ttk.Frame(order_frame)
        input_frame.grid(row=4, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        input_frame.columnconfigure(1, weight=1)
        ttk.Label(input_frame, text="Manual input overrides").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 4))
        for row, key in enumerate(PROCESS_ORDER, start=1):
            ttk.Label(input_frame, text=PROCESS_STEPS[key]["label"]).grid(row=row, column=0, sticky="w", padx=(0, 6), pady=2)
            entry = ttk.Entry(input_frame, textvariable=self.process_input_vars[key])
            entry.grid(row=row, column=1, sticky="ew", pady=2)
            browse = ttk.Button(input_frame, text="Browse", command=lambda k=key: self._browse_process_input(k))
            browse.grid(row=row, column=2, padx=(6, 0), pady=2)
            self.process_override_widgets[key] = (entry, browse)

        self._refresh_process_order_ui()

        run_frame = ttk.LabelFrame(page, text="Step 3 - Run", padding=10)
        run_frame.grid(row=4, column=0, sticky="ew", pady=(12, 0))
        self.run_process_button = ttk.Button(run_frame, text="Process images", command=self.run_processing_pipeline, state=tk.DISABLED)
        self.run_process_button.pack(side="left")
        return page

    def _create_mask_page(self, parent):
        page = ttk.Frame(parent)
        page.columnconfigure(0, weight=1)
        ttk.Label(page, text="Generate Mask", font=("Segoe UI", 20, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(page, #text="Order: 1 Select image/folder -> 2 Select UNet model -> 3 Run mask generation."
        ).grid(
            row=1, column=0, sticky="w", pady=(4, 16)
        )

        input_frame = ttk.LabelFrame(page, text="Step - 1 Input", padding=10)
        input_frame.grid(row=2, column=0, sticky="ew", pady=(0, 10))
        ttk.Button(input_frame, text="Select Image", command=self.select_mask_image).pack(side="left", padx=(0, 8))
        ttk.Button(input_frame, text="Select Folder", command=self.select_mask_folder).pack(side="left", padx=(0, 8))
        self.mask_summary = ttk.Label(input_frame, text="No input selected.")
        self.mask_summary.pack(side="left", padx=(12, 0))

        model_frame = ttk.LabelFrame(page, text="Step 2 - UNet model", padding=10)
        model_frame.grid(row=3, column=0, sticky="ew", pady=(0, 10))
        ttk.Label(model_frame, text="Model").pack(side="left")
        self.unet_model_combo = ttk.Combobox(model_frame, textvariable=self.unet_model_var, width=42, state="readonly")
        self.unet_model_combo.bind("<<ComboboxSelected>>", lambda _event: self._refresh_mask_controls())
        self.unet_model_combo.pack(side="left", padx=(6, 10))
        refresh_unet = ttk.Button(model_frame, text="Refresh", command=self.refresh_models)
        refresh_unet.pack(side="left")
        self.mask_model_widgets.extend([self.unet_model_combo, refresh_unet])

        run_frame = ttk.LabelFrame(page, text="Step 3 - Run", padding=10)
        run_frame.grid(row=4, column=0, sticky="ew", pady=6)
        self.run_mask_button = ttk.Button(run_frame, text="Generate Mask", command=self.run_mask_generation, state=tk.DISABLED)
        self.run_mask_button.pack(side="left")

        output_frame = ttk.LabelFrame(page, text="Output", padding=10)
        output_frame.grid(row=5, column=0, sticky="ew", pady=(10, 0))
        output_frame.columnconfigure(1, weight=1)
        ttk.Label(output_frame, text="Mask folder").grid(row=0, column=0, sticky="w", padx=(0, 6))
        ttk.Entry(output_frame, textvariable=self.mask_output_var, state="readonly").grid(row=0, column=1, sticky="ew")
        return page

    def _create_feature_csv_page(self, parent):
        page = ttk.Frame(parent)
        page.columnconfigure(0, weight=1)
        ttk.Label(page, text="MLP Feature CSV", font=("Segoe UI", 20, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(page).grid(row=1, column=0, sticky="w", pady=(4, 16))

        input_frame = ttk.LabelFrame(page, text="Step - 1 Input", padding=10)
        input_frame.grid(row=2, column=0, sticky="ew", pady=(0, 10))
        ttk.Button(input_frame, text="Select Image", command=self.select_image).pack(side="left", padx=(0, 8))
        ttk.Button(input_frame, text="Select Folder", command=self.select_folder).pack(side="left", padx=(0, 8))
        self.feature_summary = ttk.Label(input_frame, text="No input selected.")
        self.feature_summary.pack(side="left", padx=(12, 0))

        settings_frame = ttk.LabelFrame(page, text="Step 2 - CSV settings", padding=10)
        settings_frame.grid(row=3, column=0, sticky="ew", pady=(0, 10))
        settings_frame.columnconfigure(1, weight=1)
        ttk.Label(settings_frame, text="CSV output").grid(row=0, column=0, sticky="w", padx=(0, 6))
        feature_entry = ttk.Entry(settings_frame, textvariable=self.feature_csv_var)
        feature_entry.grid(row=0, column=1, sticky="ew")
        feature_browse = ttk.Button(settings_frame, text="Save as", command=self._browse_feature_csv_output)
        feature_browse.grid(row=0, column=2, padx=(6, 0))
        ttk.Label(settings_frame, text="Mask folder").grid(row=1, column=0, sticky="w", padx=(0, 6), pady=(6, 0))
        mask_entry = ttk.Entry(settings_frame, textvariable=self.mask_folder_var)
        mask_entry.grid(row=1, column=1, sticky="ew", pady=(6, 0))
        mask_browse = ttk.Button(settings_frame, text="Browse", command=self._browse_mask_folder)
        mask_browse.grid(row=1, column=2, padx=(6, 0), pady=(6, 0))
        self.feature_csv_widgets.extend([feature_entry, feature_browse, mask_entry, mask_browse])

        run_frame = ttk.LabelFrame(page, text="Step 3 - Run", padding=10)
        run_frame.grid(row=4, column=0, sticky="ew", pady=6)
        self.run_feature_csv_button = ttk.Button(
            run_frame,
            text="Generate Feature CSV",
            command=self.generate_feature_csv,
            state=tk.DISABLED,
        )
        self.run_feature_csv_button.pack(side="left")
        return page

    def _create_infer_page(self, parent):
        page = ttk.Frame(parent)
        page.columnconfigure(0, weight=1)
        page.rowconfigure(6, weight=1)
        ttk.Label(page, text="Model Inference & Report", font=("Segoe UI", 20, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(
            page,
            #text="Order: 1 Select image/folder -> 2 Select model and options -> 3 Generate/select feature CSV when needed -> 4 Run inference -> 5 Export CSV.",
        ).grid(row=1, column=0, sticky="w", pady=(4, 10))

        input_frame = ttk.LabelFrame(page, text="Step - 1 Input", padding=10)
        input_frame.grid(row=2, column=0, sticky="ew", pady=(6, 8))
        ttk.Button(input_frame, text="Select Image", command=self.select_image).pack(side="left", padx=(0, 8))
        ttk.Button(input_frame, text="Select Folder", command=self.select_folder).pack(side="left", padx=(0, 8))
        self.infer_summary = ttk.Label(input_frame, text="No input selected.")
        self.infer_summary.pack(side="left", padx=(12, 0))

        controls = ttk.LabelFrame(page, text="Step 2 - Model", padding=10)
        controls.grid(row=3, column=0, sticky="ew", pady=(0, 8))
        ttk.Label(controls, text="Model").pack(side="left")
        self.model_combo = ttk.Combobox(controls, textvariable=self.model_var, width=34, state="readonly")
        self.model_combo.bind("<<ComboboxSelected>>", lambda _event: self._refresh_infer_controls())
        self.model_combo.pack(side="left", padx=(6, 14))
        refresh_models_button = ttk.Button(controls, text="Refresh", command=self.refresh_models)
        refresh_models_button.pack(side="left", padx=(0, 14))
        ttk.Label(controls, text="Architecture").pack(side="left")
        arch_combo = ttk.Combobox(
            controls,
            textvariable=self.arch_var,
            values=("ResNet18", "MobileNetV2", "EfficientNetB3", "MLP-ResNet18"),
            width=18,
            state="readonly",
        )
        arch_combo.bind("<<ComboboxSelected>>", lambda _event: self._refresh_infer_controls())
        arch_combo.pack(side="left", padx=6)
        self.infer_control_widgets.extend([self.model_combo, refresh_models_button, arch_combo])

        feature_frame = ttk.LabelFrame(page, text="Step 3 - MLP feature CSV", padding=10)
        feature_frame.grid(row=4, column=0, sticky="ew", pady=(0, 8))
        feature_frame.columnconfigure(1, weight=1)
        ttk.Label(feature_frame, text="CSV").grid(row=0, column=0, sticky="w", padx=(0, 6))
        feature_entry = ttk.Entry(feature_frame, textvariable=self.feature_csv_var)
        feature_entry.grid(row=0, column=1, sticky="ew")
        feature_browse = ttk.Button(feature_frame, text="Browse", command=self._browse_feature_csv)
        feature_browse.grid(row=0, column=2, padx=(6, 0))
        self.infer_feature_widgets.extend([feature_entry, feature_browse])

        actions = ttk.LabelFrame(page, text="Step 4 / 5 - Run and export", padding=10)
        actions.grid(row=5, column=0, sticky="ew", pady=6)
        self.run_infer_button = ttk.Button(actions, text="Inference", command=self.run_inference, state=tk.DISABLED)
        self.run_infer_button.pack(side="left", padx=(0, 8))
        self.export_button = ttk.Button(actions, text="Export CSV", command=self.export_csv, state=tk.DISABLED)
        self.export_button.pack(side="left")

        columns = ("file", "class", "confidence")
        self.result_table = ttk.Treeview(page, columns=columns, show="headings")
        for col, label, width in (("file", "File", 420), ("class", "Predicted Class", 180), ("confidence", "Confidence", 120)):
            self.result_table.heading(col, text=label)
            self.result_table.column(col, width=width, anchor="w")
        self.result_table.grid(row=6, column=0, sticky="nsew", pady=(8, 0))
        return page

    def _create_settings_page(self, parent):
        page = ttk.Frame(parent)
        ttk.Label(page, text="Settings / Parameters", font=("Segoe UI", 20, "bold")).pack(anchor="w")
        ttk.Label(page, text="Not ready").pack(anchor="w", pady=20)
        ttk.Label(page, text=f"Model folder: {MODELS_DIR}").pack(anchor="w")
        ttk.Label(page, text="Execution mode: CPU").pack(anchor="w", pady=(6, 0))
        return page

    def show_page(self, name):
        for page in self.pages.values():
            page.grid_remove()
        self.pages[name].grid(row=0, column=0, sticky="nsew")
        self.current_page = name
        for page_name, button in self.nav_buttons.items():
            if page_name == name:
                button.configure(bg="#2563eb", fg="#ffffff", activebackground="#1d4ed8", activeforeground="#ffffff")
            else:
                button.configure(bg="#f3f4f6", fg="#111827", activebackground="#e5e7eb", activeforeground="#111827")

    def _set_widget_state(self, widget, enabled, readonly=False):
        if widget is None:
            return
        readonly = readonly or isinstance(widget, ttk.Combobox)
        state = "readonly" if enabled and readonly else (tk.NORMAL if enabled else tk.DISABLED)
        try:
            widget.configure(state=state)
        except tk.TclError:
            pass

    def _refresh_all_workflow_controls(self):
        self._refresh_process_order_ui()
        self._refresh_process_controls()
        self._refresh_mask_controls()
        self._refresh_feature_csv_controls()
        self._refresh_infer_controls()

    def _refresh_process_controls(self):
        has_input = self.selected_input is not None
        selected_steps = self._selected_process_steps()
        has_steps = bool(selected_steps)
        can_configure = has_input and not self.processing_busy and not self.inference_busy

        for widget in self.process_mask_widgets:
            self._set_widget_state(widget, can_configure)
        for key, widget in self.process_step_widgets.items():
            self._set_widget_state(widget, can_configure)

        list_enabled = can_configure and has_steps
        if self.process_order_list:
            self.process_order_list.configure(state=tk.NORMAL if list_enabled else tk.DISABLED)
        for widget in self.process_move_buttons:
            self._set_widget_state(widget, list_enabled)

        selected_set = set(selected_steps)
        for key, widgets in self.process_override_widgets.items():
            enabled = can_configure and has_steps and key in selected_set
            for widget in widgets:
                self._set_widget_state(widget, enabled)

        self._set_widget_state(self.run_process_button, has_input and has_steps and not self.processing_busy)

    def _refresh_mask_controls(self):
        has_input = self.selected_mask_input is not None
        has_model = bool(self.unet_model_var.get())
        can_choose_model = has_input and not self.mask_busy
        for widget in self.mask_model_widgets:
            self._set_widget_state(widget, can_choose_model, readonly=widget is self.unet_model_combo)
        self._set_widget_state(self.run_mask_button, has_input and has_model and not self.mask_busy)

    def _refresh_feature_csv_controls(self):
        has_input = self.selected_input is not None
        can_configure = has_input and not self.inference_busy
        for widget in self.feature_csv_widgets:
            self._set_widget_state(widget, can_configure)
        self._set_widget_state(self.run_feature_csv_button, can_configure)

    def _refresh_infer_controls(self):
        has_input = self.selected_input is not None
        has_model = bool(self.model_var.get())
        uses_mlp = self.arch_var.get().lower() == "mlp-resnet18"
        can_configure = has_input and not self.inference_busy

        for widget in self.infer_control_widgets:
            self._set_widget_state(widget, can_configure, readonly=widget in (self.model_combo,))
        for widget in self.infer_feature_widgets:
            self._set_widget_state(widget, can_configure and has_model and uses_mlp)

        has_feature_csv = bool(self.feature_csv_var.get().strip())
        can_run = can_configure and has_model and (not uses_mlp or has_feature_csv)
        self._set_widget_state(self.run_infer_button, can_run)
        self._set_widget_state(self.export_button, bool(self.results) and not self.inference_busy)

    def refresh_models(self):
        MODELS_DIR.mkdir(exist_ok=True)
        models = sorted(p.name for p in MODELS_DIR.glob("*.pth"))
        self.model_combo["values"] = models
        if self.unet_model_combo:
            self.unet_model_combo["values"] = models
        if models and self.model_var.get() not in models:
            self.model_var.set(models[0])
        elif not models:
            self.model_var.set("")
        if models and self.unet_model_var.get() not in models:
            self.unet_model_var.set(models[0])
        elif not models:
            self.unet_model_var.set("")
        self._refresh_all_workflow_controls()
        self.status_var.set(f"Found {len(models)} model(s) in {MODELS_DIR}")

    def select_image(self):
        paths = filedialog.askopenfilenames(filetypes=[("Images", "*.jpg *.jpeg *.png *.jfif")])
        if paths:
            self._set_input(paths)

    def select_folder(self):
        path = filedialog.askdirectory()
        if path:
            self._set_input(path)

    def select_mask_image(self):
        paths = filedialog.askopenfilenames(filetypes=[("Images", "*.jpg *.jpeg *.png *.jfif")])
        if paths:
            self._set_mask_input(paths)

    def select_mask_folder(self):
        path = filedialog.askdirectory()
        if path:
            self._set_mask_input(path)

    def _set_input(self, path):
        self.selected_input = tuple(Path(p) for p in path) if is_path_sequence(path) else Path(path)
        count = len(list_images(self.selected_input))
        text = self._selection_text(count, self.selected_input)
        self.process_summary.config(text=text)
        self.infer_summary.config(text=text)
        if self.feature_summary:
            self.feature_summary.config(text=text)
        self.status_var.set(text)
        self._refresh_all_workflow_controls()

    def _set_mask_input(self, path):
        self.selected_mask_input = tuple(Path(p) for p in path) if is_path_sequence(path) else Path(path)
        count = len(list_images(self.selected_mask_input))
        text = self._selection_text(count, self.selected_mask_input)
        self.mask_summary.config(text=text)
        self.status_var.set(text)
        self._refresh_mask_controls()

    def _selection_text(self, count, selection):
        if is_path_sequence(selection):
            base = result_parent_for(selection)
            return f"Selected {count} image(s) from multiple files under: {base}"
        return f"Selected {count} image(s): {selection}"

    def _selected_process_steps(self):
        return [key for key in self.custom_process_order if self.process_step_vars[key].get()]

    def _refresh_process_order_ui(self):
        if not self.process_order_list:
            return
        self.process_order_list.configure(state=tk.NORMAL)
        selected = self._selected_process_steps()
        self.process_order_list.delete(0, tk.END)
        for key in selected:
            self.process_order_list.insert(tk.END, PROCESS_STEPS[key]["label"])
        self._refresh_process_controls()

    def _move_process_step(self, direction):
        if not self.process_order_list:
            return
        selection = self.process_order_list.curselection()
        if not selection:
            return
        selected_steps = self._selected_process_steps()
        index = selection[0]
        new_index = index + direction
        if new_index < 0 or new_index >= len(selected_steps):
            return
        selected_steps[index], selected_steps[new_index] = selected_steps[new_index], selected_steps[index]
        unselected = [key for key in self.custom_process_order if key not in selected_steps]
        self.custom_process_order = selected_steps + unselected
        self._refresh_process_order_ui()
        self.process_order_list.selection_set(new_index)

    def _browse_process_input(self, key):
        path = filedialog.askdirectory(title=f"Select input folder for {PROCESS_STEPS[key]['label']}")
        if path:
            self.process_input_vars[key].set(path)

    def _browse_process_mask_folder(self):
        path = filedialog.askdirectory(title="Select mask folder for image classification split")
        if path:
            self.process_mask_folder_var.set(path)

    def _clear_process_mask_folder(self):
        self.process_mask_folder_var.set("")

    def _browse_feature_csv(self):
        path = filedialog.askopenfilename(filetypes=[("CSV", "*.csv")])
        if path:
            self.feature_csv_var.set(path)

    def _browse_feature_csv_output(self):
        path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            initialfile="mlp_features.csv",
            filetypes=[("CSV", "*.csv")],
        )
        if path:
            self.feature_csv_var.set(path)

    def _browse_mask_folder(self):
        path = filedialog.askdirectory(title="Select mask folder")
        if path:
            self.mask_folder_var.set(path)

    def generate_feature_csv(self):
        images = list_images(self.selected_input) if self.selected_input else []
        if not images:
            messagebox.showwarning("No images", "Select an image or folder first.")
            return
        path = self.feature_csv_var.get().strip()
        if not path:
            path = filedialog.asksaveasfilename(
                defaultextension=".csv",
                initialfile="mlp_features.csv",
                filetypes=[("CSV", "*.csv")],
            )
            if not path:
                return
        self.feature_csv_var.set(path)
        self._run_background(lambda: self._feature_csv_worker(images, Path(path)))

    def _feature_csv_worker(self, images, output_csv):
        try:
            mask_dir = Path(self.mask_folder_var.get()) if self.mask_folder_var.get().strip() else None
            base_metadata = FEATURES_DIR / "metadata.csv"
            csv_path, row_count, missing_masks = generate_feature_csv(
                images,
                output_csv,
                mask_dir=mask_dir,
                base_metadata_csv=base_metadata if base_metadata.exists() else None,
                progress=self._progress,
            )
            self._done(f"Feature CSV exported: {csv_path} ({row_count} rows, {missing_masks} missing masks)")
        except Exception as exc:
            traceback.print_exc()
            self.root.after(0, lambda: messagebox.showerror("Feature CSV failed", str(exc)))
            self._done("Feature CSV failed.")

    def _default_feature_csv_path(self, images=None):
        if images:
            root = image_root_for(images[0])
        else:
            root = image_root_for(self.selected_input) if self.selected_input else RUNTIME_DIR
        return root / "mlp_features.csv"

    def _prepare_mlp_metadata(self, images):
        stats = load_train_stats()
        csv_text = self.feature_csv_var.get().strip()
        if not csv_text:
            raise ValueError("MLP-ResNet18 inference needs a feature CSV. Generate or select it from the Feature CSV page first.")
        csv_path = Path(csv_text)
        if not csv_path.exists():
            raise ValueError(f"Feature CSV does not exist: {csv_path}")
        return build_metadata_tensors(csv_path, stats), stats, csv_path

    def _confirm_process_steps(self, steps):
        if not steps:
            messagebox.showwarning("No processing steps", "Select at least one processing step.")
            return False
        for key in steps:
            warning = PROCESS_STEPS[key].get("warning")
            if warning and not messagebox.askyesno(PROCESS_STEPS[key]["label"], f"{warning}\n\nContinue?"):
                return False
        warning = self._process_order_warning(steps)
        if warning and not messagebox.askyesno("Processing order warning", f"{warning}\n\nContinue anyway?"):
            return False
        return True

    def _process_order_warning(self, steps):
        positions = {key: index for index, key in enumerate(steps)}
        warnings = []
        has_crop_before_rb = "crop" in positions and positions.get("crop", 99) < positions.get("rb", -1)
        has_source_before_assort = any(
            key in positions and positions[key] < positions.get("assort", -1)
            for key in ("crop", "rb")
        )
        if "rb" in positions and not has_crop_before_rb and not self.process_input_vars["rb"].get().strip():
            warnings.append("RB usually expects IMG_Cut. Crop is not before RB and no RB input override is set.")
        if "assort" in positions and not has_source_before_assort and not self.process_input_vars["assort"].get().strip():
            warnings.append("Image classification usually expects IMG_RB or IMG_Cut. No earlier source step or input override is set.")
        return "\n".join(warnings)

    def run_processing_pipeline(self):
        if not self.selected_input:
            messagebox.showwarning("No input", "Select an image or folder first.")
            return
        steps = self._selected_process_steps()
        if not self._confirm_process_steps(steps):
            return
        self.processing_busy = True
        self._refresh_all_workflow_controls()
        self._run_background(lambda: self._processing_worker(steps))

    def _processing_worker(self, steps):
        try:
            final_source = self._execute_processing_pipeline(steps)
            self._done(f"Processing complete. Current image source: {final_source}")
        except Exception as exc:
            traceback.print_exc()
            self.root.after(0, lambda: messagebox.showerror("Processing failed", str(exc)))
            self._done("Processing failed.")
        finally:
            self.root.after(0, self._finish_processing_busy)

    def _finish_processing_busy(self):
        self.processing_busy = False
        self._refresh_all_workflow_controls()

    def _execute_processing_pipeline(self, steps):
        run_dir = make_result_run(self.selected_input)
        self.last_result_run = run_dir
        base_input, _copied = copy_images_to_workspace(self.selected_input, run_dir / "source")
        current_source = base_input
        outputs = {}
        override_sources = {}
        total = len(steps)
        for index, key in enumerate(steps, 1):
            step_input = self._resolve_process_input(key, base_input, current_source, outputs, run_dir, override_sources)
            self._progress(index - 1, total, f"Running {PROCESS_STEPS[key]['label']} on {step_input}")
            step_output = self._execute_process_step(key, step_input)
            if step_output:
                outputs[key] = step_output
                current_source = step_output
            self._progress(index, total, f"Finished {PROCESS_STEPS[key]['label']}")
        return current_source

    def _resolve_process_input(self, key, base_input, current_source, outputs, run_dir, override_sources):
        override = self.process_input_vars[key].get().strip()
        if override:
            if key not in override_sources:
                override_sources[key], _copied = copy_images_to_workspace(Path(override), run_dir / f"input_{key}")
            return override_sources[key]
        if key == "rb" and "crop" in outputs:
            return outputs["crop"]
        if key == "assort":
            if "rb" in outputs and list_images(outputs["rb"]):
                return outputs["rb"]
            if "crop" in outputs and list_images(outputs["crop"]):
                return outputs["crop"]
        return current_source

    def _execute_process_step(self, key, step_input):
        if not step_input.exists() or not step_input.is_dir():
            raise ValueError(f"Input folder does not exist for {PROCESS_STEPS[key]['label']}: {step_input}")
        ensure_stdio_for_windowed_app()

        def callback(*_args):
            self.root.after(0, lambda: self.status_var.set(f"Running {PROCESS_STEPS[key]['label']}..."))

        if key == "convert":
            from photo_process.format_converter import process_folder
            process_folder(step_input, callback)
            return step_input
        if key == "crop":
            from photo_process.batch_crop import process_folder
            process_folder(step_input, callback)
            return step_input.parent / "IMG_Cut"
        if key == "rb":
            from photo_process.image_RB import process_folder
            output_dir = step_input.parent / "IMG_RB"
            process_folder(step_input, output_dir, callback)
            return output_dir
        if key == "assort":
            from photo_process.auto_assort import process_folder
            mask_folder = self.process_mask_folder_var.get().strip()
            process_folder(step_input, callback, mask_folder=Path(mask_folder) if mask_folder else None)
            return step_input.parent / "IMG_train"
        raise ValueError(f"Unknown processing step: {key}")

    def run_mask_generation(self):
        if not self.selected_mask_input:
            messagebox.showwarning("No input", "Select an image or folder first.")
            return
        if not self.unet_model_var.get():
            messagebox.showwarning("No UNet model", "Select a UNet .pth file from the models folder.")
            return
        model_path = MODELS_DIR / self.unet_model_var.get()
        self.mask_busy = True
        self._refresh_mask_controls()
        self._run_background(lambda: self._mask_generation_worker(model_path))

    def _mask_generation_worker(self, model_path):
        try:
            run_dir = make_result_run(self.selected_mask_input)
            self.last_result_run = run_dir
            source_dir, images = copy_images_to_workspace(self.selected_mask_input, run_dir / "source")
            output_dir = run_dir / "IMG_Mask"
            unet_dir = APP_DIR / "unet"
            if str(unet_dir) not in sys.path:
                sys.path.insert(0, str(unet_dir))
            from unet.mask_generator import generate_masks

            def progress(done, total, text):
                self._progress(done, total, text)

            count = generate_masks(
                images,
                source_dir,
                output_dir,
                model_path,
                progress=progress,
            )
            self.root.after(0, lambda: self.mask_output_var.set(str(output_dir)))
            self.root.after(0, lambda: self.mask_folder_var.set(str(output_dir)))
            self._done(f"Generated {count} mask(s): {output_dir}")
        except Exception as exc:
            traceback.print_exc()
            self.root.after(0, lambda: messagebox.showerror("Mask generation failed", str(exc)))
            self._done("Mask generation failed.")
        finally:
            self.root.after(0, self._finish_mask_busy)

    def _finish_mask_busy(self):
        self.mask_busy = False
        self._refresh_mask_controls()

    def run_inference(self):
        images = list_images(self.selected_input) if self.selected_input else []
        if not images:
            messagebox.showwarning("No images", "Select an image or folder first.")
            return
        if not self.model_var.get():
            messagebox.showwarning("No model", "Put .pth files in the models folder and refresh.")
            return
        if self.arch_var.get().lower() == "mlp-resnet18" and not self.feature_csv_var.get().strip():
            messagebox.showwarning("No feature CSV", "Generate or select a feature CSV from the Feature CSV page first.")
            return
        self.results = []
        self._refresh_infer_controls()
        for item in self.result_table.get_children():
            self.result_table.delete(item)
        model_path = MODELS_DIR / self.model_var.get()
        arch = self.arch_var.get()
        self.inference_busy = True
        self._refresh_all_workflow_controls()
        self._run_background(lambda: self._inference_worker(images, model_path, arch))

    def _inference_worker(self, images, model_path, arch):
        try:
            metadata = None
            stats = None
            metadata_csv = None
            if arch.lower() == "mlp-resnet18":
                metadata, stats, metadata_csv = self._prepare_mlp_metadata(images)
            engine = InferenceEngine(model_path, arch, metadata=metadata, stats=stats)
            for index, image_path in enumerate(images, 1):
                pred_index, label, confidence, probs = engine.predict(image_path)
                row = {
                    "file": str(image_path),
                    "predicted_class": label,
                    "predicted_index": pred_index,
                    "confidence": confidence,
                    "probabilities": probs,
                }
                if metadata_csv:
                    row["metadata_csv"] = str(metadata_csv)
                self.results.append(row)
                self.root.after(0, lambda r=row: self.result_table.insert("", "end", values=(Path(r["file"]).name, r["predicted_class"], f"{r['confidence']:.4f}")))
                self._progress(index, len(images), f"Predicted {index}/{len(images)}")
            self._done(self._summary_text())
        except Exception as exc:
            traceback.print_exc()
            self.root.after(0, lambda: messagebox.showerror("Inference failed", str(exc)))
            self._done("Inference failed.")
        finally:
            self.root.after(0, self._finish_inference_busy)

    def _finish_inference_busy(self):
        self.inference_busy = False
        self._refresh_all_workflow_controls()

    def _summary_text(self):
        if not self.results:
            return "No results."
        counts = Counter(r["predicted_class"] for r in self.results)
        averages = defaultdict(list)
        for row in self.results:
            averages[row["predicted_class"]].append(row["confidence"])
        parts = [f"{name}: {counts[name]} avg {sum(averages[name]) / len(averages[name]):.3f}" for name in sorted(counts)]
        return f"Completed {len(self.results)} prediction(s). " + "; ".join(parts)

    def export_csv(self):
        if not self.results:
            messagebox.showwarning("No results", "Run inference before exporting a report.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV", "*.csv")])
        if not path:
            return
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=["file", "predicted_index", "predicted_class", "confidence"])
            writer.writeheader()
            for row in self.results:
                writer.writerow({k: row[k] for k in writer.fieldnames})
        self.status_var.set(f"CSV report exported: {path}")

    def _run_background(self, target):
        self.progress_var.set(0)
        threading.Thread(target=target, daemon=True).start()

    def _progress(self, current, total, text):
        percent = (current / total) * 100 if total else 0
        self.root.after(0, lambda: (self.progress_var.set(percent), self.status_var.set(text)))

    def _done(self, text):
        self.root.after(0, lambda: (self.progress_var.set(100), self.status_var.set(text)))


def main():
    root = tk.Tk()
    DeployApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
