"""GUI-free feature extraction with the legacy CSV contract."""

from __future__ import annotations

import csv
import json
import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping

from .contracts import FEATURE_CSV_FIELDS
from .models import MLDependencyError


IMAGE_SUFFIXES = frozenset({".jpg", ".jpeg", ".png", ".bmp", ".webp", ".jfif"})


def safe_float(value: Any, default: float | None = 0.0) -> float | None:
    if value in (None, ""):
        return default
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default
    return default if math.isnan(number) else number


def read_csv_rows(path: str | Path) -> list[dict[str, str]]:
    path = Path(path)
    if not path.exists():
        return []
    for encoding in ("utf-8-sig", "utf-8", "cp950"):
        try:
            with path.open("r", newline="", encoding=encoding) as handle:
                return list(csv.DictReader(handle))
        except UnicodeDecodeError:
            continue
    with path.open("r", newline="", encoding="utf-8", errors="replace") as handle:
        return list(csv.DictReader(handle))


def metadata_by_name(path: str | Path | None) -> dict[str, dict[str, str]]:
    if not path:
        return {}
    result: dict[str, dict[str, str]] = {}
    for row in read_csv_rows(path):
        name = row.get("image_name") or row.get("file") or row.get("filename")
        if name:
            result[Path(name).name] = row
    return result


def load_train_stats(path: str | Path | None) -> dict[str, dict[str, float]]:
    from .contracts import DEFAULT_TRAIN_STATS, validate_train_stats

    if path and Path(path).exists():
        with Path(path).open("r", encoding="utf-8-sig") as handle:
            return validate_train_stats(json.load(handle))
    return validate_train_stats(DEFAULT_TRAIN_STATS)


def index_masks(folder: str | Path | None) -> dict[str, Path]:
    if not folder:
        return {}
    folder = Path(folder)
    if not folder.exists():
        return {}
    result: dict[str, Path] = {}
    for path in sorted(folder.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in IMAGE_SUFFIXES:
            continue
        result[path.stem] = path
        if path.stem.endswith("_mask"):
            result[path.stem[:-5]] = path
    return result


def calculate_image_features(image_path: str | Path, mask_path: str | Path) -> dict[str, float]:
    """Calculate Length_px, Entropy and Contrast_Ratio exactly as app.py."""

    try:
        import cv2
        import numpy as np
        from scipy.ndimage import find_objects
        from skimage import measure
        from skimage.feature import graycomatrix
    except ImportError as exc:
        raise MLDependencyError(
            "Feature extraction requires NumPy, OpenCV, SciPy and scikit-image."
        ) from exc

    config = {
        "glcm_levels": 32,
        "glcm_distances": [1, 2],
        "glcm_angles": [0, np.pi / 4, np.pi / 2, 3 * np.pi / 4],
        "contrast_kernel_min": 20,
        "contrast_kernel_max": 50,
        "contrast_kernel_scale": 0.35,
        "min_roi_pixels": 100,
        "denoise_kernel": 3,
        "epsilon": 1e-10,
    }
    image_path = Path(image_path)
    mask_path = Path(mask_path)
    image = cv2.imdecode(np.fromfile(str(image_path), dtype=np.uint8), cv2.IMREAD_COLOR)
    mask = cv2.imdecode(np.fromfile(str(mask_path), dtype=np.uint8), cv2.IMREAD_GRAYSCALE)
    if image is None or mask is None:
        raise ValueError(f"Cannot decode image or mask: {image_path.name}")
    return calculate_features_array(image, mask)


def calculate_features_array(image, mask):
    import cv2
    import numpy as np
    from scipy.ndimage import find_objects
    from skimage import measure
    from skimage.feature import graycomatrix

    config = {
        "glcm_levels": 32,
        "glcm_distances": [1, 2],
        "glcm_angles": [0, np.pi / 4, np.pi / 2, 3 * np.pi / 4],
        "contrast_kernel_min": 20,
        "contrast_kernel_max": 50,
        "contrast_kernel_scale": 0.35,
        "min_roi_pixels": 100,
        "denoise_kernel": 3,
        "epsilon": 1e-10,
    }
    if mask.shape != image.shape[:2]:
        mask = cv2.resize(mask, (image.shape[1], image.shape[0]), interpolation=cv2.INTER_NEAREST)
    _, mask_binary = cv2.threshold(mask, 127, 1, cv2.THRESH_BINARY)
    if int(np.sum(mask_binary)) < config["min_roi_pixels"]:
        return {"Length_px": 0.0, "Entropy": 0.0, "Contrast_Ratio": 0.0}

    labeled = measure.label(mask_binary)
    regions = measure.regionprops(labeled)
    region = max(regions, key=lambda candidate: candidate.area) if regions else None
    length = float(region.major_axis_length) if region else 0.0

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    if config["denoise_kernel"] > 0:
        gray = cv2.medianBlur(gray, config["denoise_kernel"])
    roi_pixels = gray[mask_binary == 1]
    mean_intensity = float(np.mean(roi_pixels)) if roi_pixels.size else 0.0

    entropy = 0.0
    if find_objects(labeled) and region:
        top, left, bottom, right = region.bbox
        cropped_gray = gray[top:bottom, left:right]
        cropped_mask = mask_binary[top:bottom, left:right]
        masked_gray = cropped_gray.copy()
        masked_gray[cropped_mask == 0] = 0
        maximum = np.max(masked_gray)
        if maximum > 0:
            levels = config["glcm_levels"]
            scaled = (masked_gray / maximum * (levels - 1)).astype(np.uint8)
            glcm = graycomatrix(
                scaled,
                distances=config["glcm_distances"],
                angles=config["glcm_angles"],
                levels=levels,
                symmetric=True,
                normed=True,
            )
            probabilities = glcm / (glcm.sum(axis=(0, 1), keepdims=True) + config["epsilon"])
            entropy = float(
                np.mean(
                    -np.sum(
                        probabilities * np.log2(probabilities + config["epsilon"]),
                        axis=(0, 1),
                    )
                )
            )

    roi_area = np.sum(mask_binary)
    kernel_size = int(np.sqrt(roi_area) * config["contrast_kernel_scale"])
    kernel_size = int(
        np.clip(kernel_size, config["contrast_kernel_min"], config["contrast_kernel_max"])
    )
    kernel_size = kernel_size if kernel_size % 2 else kernel_size + 1
    kernel = np.ones((kernel_size, kernel_size), np.uint8)
    dilated = cv2.dilate(mask_binary.astype(np.uint8), kernel, iterations=1)
    ring = np.clip(dilated - mask_binary, 0, 1)
    background = gray[ring == 1]
    contrast = (
        1.0
        if background.size < 10
        else float(mean_intensity / (np.mean(background) + config["epsilon"]))
    )
    return {"Length_px": length, "Entropy": entropy, "Contrast_Ratio": contrast}


def load_diabetes_reference(excel_path: str | Path | None) -> list[tuple[float, int]]:
    if not excel_path or not Path(excel_path).exists():
        return []
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise MLDependencyError("openpyxl is required to read the diabetes reference.") from exc
    workbook = load_workbook(Path(excel_path), data_only=True, read_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    egfr_index = headers.index("matched_eGFR") if "matched_eGFR" in headers else 0
    diabetes_index = next(
        (
            index
            for index, header in enumerate(headers)
            if "diabetes" in header.lower() or header.lower() == "dm" or "糖尿" in header
        ),
        None,
    )
    if diabetes_index is None:
        return []
    result: list[tuple[float, int]] = []
    for row in rows[1:]:
        if len(row) <= max(egfr_index, diabetes_index):
            continue
        egfr = safe_float(row[egfr_index], None)
        if egfr is not None:
            result.append((egfr, int(safe_float(row[diabetes_index], 0.0) or 0)))
    return result


def diabetes_from_filename(
    image_name: str,
    reference: Iterable[tuple[float, int]],
    tolerance: float = 0.2,
) -> int:
    try:
        egfr = float(Path(image_name).stem.split("_")[0])
    except (IndexError, ValueError):
        return 0
    reference = list(reference)
    if not reference:
        return 0
    closest = min(reference, key=lambda item: abs(item[0] - egfr))
    return int(closest[1]) if abs(closest[0] - egfr) <= tolerance else 0


@dataclass
class FeatureReport:
    output_csv: Path
    rows: int
    missing_masks: list[str] = field(default_factory=list)
    missing_metadata: list[str] = field(default_factory=list)
    failed_features: dict[str, str] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "output_csv": self.output_csv.name,
            "rows": self.rows,
            "missing_masks": self.missing_masks,
            "missing_metadata": self.missing_metadata,
            "failed_features": self.failed_features,
        }


def write_feature_csv(path: str | Path, rows: Iterable[Mapping[str, Any]]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=FEATURE_CSV_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in FEATURE_CSV_FIELDS})
    return path


def generate_feature_csv(
    image_paths: Iterable[str | Path],
    output_csv: str | Path,
    *,
    mask_dir: str | Path | None,
    metadata_csv: str | Path | None = None,
    diabetes_excel: str | Path | None = None,
    progress: Callable[[int, int, str], None] | None = None,
) -> FeatureReport:
    images = sorted(Path(item) for item in image_paths)
    if not images:
        raise ValueError("No images supplied for feature extraction.")
    masks = index_masks(mask_dir)
    metadata = metadata_by_name(metadata_csv)
    diabetes_reference = load_diabetes_reference(diabetes_excel)
    report = FeatureReport(output_csv=Path(output_csv), rows=0)
    rows: list[dict[str, Any]] = []
    for index, image_path in enumerate(images, 1):
        base = metadata.get(image_path.name)
        if base is None:
            base = {}
            report.missing_metadata.append(image_path.name)
        row = {field: "" for field in FEATURE_CSV_FIELDS}
        row.update({field: base.get(field, "") for field in FEATURE_CSV_FIELDS})
        row["image_name"] = image_path.name
        if row["has_diabetes"] == "":
            row["has_diabetes"] = diabetes_from_filename(image_path.name, diabetes_reference)
        mask = masks.get(image_path.stem)
        if mask is None:
            report.missing_masks.append(image_path.name)
        else:
            try:
                row.update(calculate_image_features(image_path, mask))
            except Exception as exc:
                report.failed_features[image_path.name] = str(exc)
        rows.append(row)
        if progress:
            progress(index, len(images), f"feature {index}/{len(images)}")
    write_feature_csv(output_csv, rows)
    report.rows = len(rows)
    return report
